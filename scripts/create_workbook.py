"""Build the Stock Analyzer workbook (templates/stocks.template.xlsx, or stocks.xlsx).

    python scripts/create_workbook.py --template    # rebuild the committed template
    python scripts/create_workbook.py               # create stocks.xlsx if it doesn't exist
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from stock_analyzer import config
from stock_analyzer.excel_io import LOG_COLUMNS, SCORECARD_SHEET, STOCK_COLUMNS

MAX_ROWS = 500
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
AUTO_FILL = PatternFill("solid", fgColor="6B7280")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RATING_COLORS = {
    "Buy": "C6EFCE", "Overweight": "E2F0D9", "Hold": "FFF2CC", "Underweight": "FCE4D6", "Sell": "F8CBAD",
    "REVIEW": "D9D9D9",
}

STOCK_HELP = {
    "Ticker": "Exchange symbol without suffix, e.g. RELIANCE, TCS, AAPL, MSFT.",
    "Exchange": "NSE, BSE, NASDAQ or NYSE.",
    "Company": "Optional. Filled in automatically on the first run.",
    "Status": "Holding = you own it. Watchlist = you are considering it.",
    "Investment Horizon": "How long you plan to hold. Shorter horizons weight technicals; longer ones weight fundamentals.",
    "Qty": "Shares held (holdings only).",
    "Avg Buy Price": "Your average cost per share, in the stock's currency.",
    "Buy Date": "When you first bought (optional).",
    "Target Return %": "Optional return you are aiming for over the horizon, e.g. 20.",
    "Risk Tolerance": "Low, Med or High. Shapes stop distances and position advice.",
    "Depth": "Quick (~5 LLM calls), Standard (~11) or Deep (~16). Blank = Settings default.",
    "2nd Opinion": "Y = also run TradingAgents (needs ANTHROPIC_API_KEY).",
    "Analyze?": "N skips this stock without deleting the row.",
    "My Thesis / Notes": "Your own view. The agent reads it and will challenge it.",
    "Last Analyzed": "Written by the agent.",
    "Last Rating": "Written by the agent: Buy / Overweight / Hold / Underweight / Sell / REVIEW.",
    "Last Action": "Written by the agent: what the rating means for your position.",
    "Last Report": "Written by the agent: link to the latest HTML report.",
}

SETTINGS_HELP = {
    "Data Provider": "OpenBB provider: yfinance (free, India + US) or fmp (needs FMP_API_KEY and openbb-fmp).",
    "Stage Model": "Claude model for analyst, persona, debate and risk stages (run.py only).",
    "Decision Model": "Claude model for the research manager and portfolio manager (run.py only).",
    "Default Depth": "Quick, Standard or Deep, used when a stock's Depth is blank.",
    "Debate Rounds": "Bull vs bear rounds at Deep depth (Standard always uses 1).",
    "Personas (short)": "Investor personas for horizons up to 1 month (prompts/personas/<name>.md).",
    "Personas (medium)": "Personas for 3-6 month horizons.",
    "Personas (long)": "Personas for 1 year and longer.",
    "Settle Days (short)": "Trading days after which a short-horizon call is scored against what happened.",
    "Settle Days (medium)": "Same, for medium horizons.",
    "Settle Days (long)": "Same, for long horizons.",
    "Use Kite": "Y = Claude Code agent may read live Indian quotes from a Zerodha Kite connector, if logged in.",
    "Use Web Search": "Y = agents may search the web for news and sentiment.",
    "TradingAgents Quick Model": "Model for TradingAgents' quick-thinking agents (2nd opinion).",
    "TradingAgents Deep Model": "Model for TradingAgents' deep-thinking managers (2nd opinion).",
}

HOW_TO = [
    ("Stock Analyzer: how to use this workbook", True),
    ("", False),
    ("1. Stocks sheet: one row per stock. Fill the white columns; the grey ones are written by the agent.", False),
    ("2. Keep this file CLOSED while an analysis runs. The agent will not write to it while Excel has it open.", False),
    ("3. Run in Claude Code:  /analyze-stocks   (all stocks)   /analyze-stocks RELIANCE   (one)", False),
    ("   or standalone:       .venv/bin/python run.py --ticker RELIANCE", False),
    ("4. Reports: reports/<date>/index.html lists the day; each stock has reports/<date>/<TICKER_EXCHANGE>/report.html.", False),
    ("5. Analysis Log: one row per stock per run. After the settle window the agent fills Return, Alpha, Outcome and a Lesson.", False),
    ("6. Scorecard: how often each rating turned out right. Rebuilt on every run.", False),
    ("7. Settings: models, depth, personas and data provider.", False),
    ("", False),
    ("Ratings: Buy > Overweight > Hold > Underweight > Sell. REVIEW means the run could not reach a clear call.", False),
    ("Holding actions: Add / Add on dips / Hold / Trim / Exit.  Watchlist actions: Buy / Accumulate / Wait / Avoid.", False),
    ("", False),
    ("This is research output, not investment advice.", True),
]

EXAMPLES = [
    {"Ticker": "RELIANCE", "Exchange": "NSE", "Status": "Holding", "Investment Horizon": "1 Year", "Qty": 50,
     "Avg Buy Price": 1250, "Buy Date": date(2025, 11, 3), "Target Return %": 15, "Risk Tolerance": "Med",
     "Depth": "", "2nd Opinion": "N", "Analyze?": "Y",
     "My Thesis / Notes": "Example row. Jio and retail growth should offset weak refining margins."},
    {"Ticker": "AAPL", "Exchange": "NASDAQ", "Status": "Watchlist", "Investment Horizon": "3 Months",
     "Target Return %": 10, "Risk Tolerance": "Low", "Depth": "", "2nd Opinion": "N", "Analyze?": "Y",
     "My Thesis / Notes": "Example row. Looking for an entry after the iPhone cycle news settles."},
]


def _header(ws, columns, auto=()):
    for i, name in enumerate(columns, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = HEADER_FONT
        c.fill = AUTO_FILL if name in auto else HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32


def _list_validation(ws, col_idx, values, prompt):
    letter = get_column_letter(col_idx)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error, dv.errorTitle = f"Choose one of: {', '.join(values)}", "Invalid value"
    dv.prompt, dv.promptTitle = prompt, ws.cell(row=1, column=col_idx).value
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{MAX_ROWS}")


def _rating_colors(ws, col_idx):
    letter = get_column_letter(col_idx)
    for rating, color in RATING_COLORS.items():
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{MAX_ROWS}",
            CellIsRule(operator="equal", formula=[f'"{rating}"'], fill=PatternFill("solid", fgColor=color)),
        )


def build(path: Path, examples: bool = True) -> Path:
    wb = Workbook()

    ws = wb.active
    ws.title = "Stocks"
    auto = {"Last Analyzed", "Last Rating", "Last Action", "Last Report"}
    _header(ws, STOCK_COLUMNS, auto)
    widths = {"Ticker": 12, "Exchange": 10, "Company": 26, "Status": 11, "Investment Horizon": 13, "Qty": 8,
              "Avg Buy Price": 11, "Buy Date": 11, "Target Return %": 9, "Risk Tolerance": 9, "Depth": 10,
              "2nd Opinion": 9, "Analyze?": 9, "My Thesis / Notes": 48, "Last Analyzed": 12,
              "Last Rating": 12, "Last Action": 14, "Last Report": 14}
    for i, name in enumerate(STOCK_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 12)
        ws.cell(row=1, column=i).comment = Comment(STOCK_HELP[name], "Stock Analyzer")
    col = {n: i for i, n in enumerate(STOCK_COLUMNS, 1)}
    _list_validation(ws, col["Exchange"], list(config.EXCHANGES), "NSE / BSE / NASDAQ / NYSE")
    _list_validation(ws, col["Status"], ["Holding", "Watchlist"], "Do you own it?")
    _list_validation(ws, col["Investment Horizon"], config.HORIZONS, "How long will you hold?")
    _list_validation(ws, col["Risk Tolerance"], ["Low", "Med", "High"], "Your risk appetite for this stock")
    _list_validation(ws, col["Depth"], config.DEPTHS, "Blank = Settings default")
    _list_validation(ws, col["2nd Opinion"], ["Y", "N"], "Also run TradingAgents?")
    _list_validation(ws, col["Analyze?"], ["Y", "N"], "Include in the daily run?")
    _rating_colors(ws, col["Last Rating"])
    for c in ("Buy Date", "Last Analyzed"):
        for r in range(2, MAX_ROWS + 1):
            ws.cell(row=r, column=col[c]).number_format = "yyyy-mm-dd"
    if examples:
        for r, ex in enumerate(EXAMPLES, 2):
            for name, value in ex.items():
                ws.cell(row=r, column=col[name], value=value if value != "" else None)

    log = wb.create_sheet("Analysis Log")
    _header(log, LOG_COLUMNS, {"Settled On", "Return", "Benchmark Return", "Alpha", "Outcome", "Lesson"})
    for i, name in enumerate(LOG_COLUMNS, 1):
        log.column_dimensions[get_column_letter(i)].width = {"Lesson": 60, "Summary": 50, "Persona Votes": 30, "Report": 14,
                                                              "Entry Zone": 16}.get(name, 11)
    _rating_colors(log, LOG_COLUMNS.index("Rating") + 1)
    _rating_colors(log, LOG_COLUMNS.index("2nd Opinion") + 1)

    sc = wb.create_sheet(SCORECARD_SHEET)
    sc["A1"] = "Scorecard is rebuilt after every run, once calls have been settled."
    sc["A1"].font = Font(italic=True)

    st = wb.create_sheet("Settings")
    _header(st, ["Setting", "Value", "What it does"])
    st.column_dimensions["A"].width, st.column_dimensions["B"].width, st.column_dimensions["C"].width = 26, 26, 90
    for r, (k, v) in enumerate(config.DEFAULT_SETTINGS.items(), 2):
        st.cell(row=r, column=1, value=k).font = Font(bold=True)
        st.cell(row=r, column=2, value=v)
        st.cell(row=r, column=3, value=SETTINGS_HELP.get(k, ""))

    how = wb.create_sheet("How To Use")
    how.column_dimensions["A"].width = 120
    for r, (line, bold) in enumerate(HOW_TO, 1):
        how.cell(row=r, column=1, value=line).font = Font(bold=bold, size=13 if r == 1 else 11)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", action="store_true", help="(re)build templates/stocks.template.xlsx")
    ap.add_argument("--force", action="store_true", help="overwrite stocks.xlsx if it exists")
    args = ap.parse_args()
    if args.template:
        print(build(config.TEMPLATE_WORKBOOK))
        return
    if config.WORKBOOK.exists() and not args.force:
        sys.exit(f"{config.WORKBOOK} already exists; pass --force to overwrite it (your data would be lost).")
    print(build(config.WORKBOOK))


if __name__ == "__main__":
    main()
