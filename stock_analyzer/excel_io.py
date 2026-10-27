"""Read and write stocks.xlsx.

Columns are found by header name, so you can reorder or add your own columns
without breaking anything. Every write first checks that Excel doesn't have the
file open, so your unsaved edits are never overwritten.
"""

from __future__ import annotations

import os
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

from . import config

STOCK_COLUMNS = [
    "Ticker", "Exchange", "Company", "Status", "Investment Horizon", "Qty", "Avg Buy Price", "Buy Date",
    "Target Return %", "Risk Tolerance", "Depth", "2nd Opinion", "Analyze?", "My Thesis / Notes",
    "Last Analyzed", "Last Rating", "Last Action", "Last Report",
]

LOG_COLUMNS = [
    "Date", "Ticker", "Exchange", "Horizon", "Status", "Depth", "Data As Of", "Price", "Currency",
    "Technical", "Fundamental", "News", "Sentiment", "Composite", "Persona Votes", "Research Plan",
    "Rating", "Action", "Confidence", "Summary", "Entry Zone", "Target", "Stop", "2nd Opinion", "Report",
    "Settled On", "Return", "Benchmark Return", "Alpha", "Outcome", "Lesson",
]

SCORECARD_SHEET = "Scorecard"


class WorkbookOpenError(RuntimeError):
    pass


@dataclass
class StockRow:
    row: int
    ticker: str
    exchange: str
    company: str | None
    status: str
    horizon: str
    qty: float | None
    avg_price: float | None
    buy_date: str | None
    target_return: float | None
    risk_tolerance: str
    depth: str | None
    second_opinion: bool
    analyze: bool
    notes: str | None
    last_analyzed: str | None = None
    errors: list[str] = field(default_factory=list)

    @property
    def key(self) -> str:
        return config.stock_key(self.ticker, self.exchange)

    @property
    def market(self) -> str:
        return config.exchange(self.exchange).market

    @property
    def bucket(self) -> str:
        return config.horizon_bucket(self.horizon)

    @property
    def holding(self) -> bool:
        return self.status.lower() == "holding"

    @property
    def position(self) -> dict | None:
        if not self.holding or not self.qty or not self.avg_price:
            return None
        return {"qty": self.qty, "avg_price": self.avg_price, "buy_date": self.buy_date}

    def as_dict(self) -> dict:
        d = {k: v for k, v in self.__dict__.items() if k != "errors"}
        d.update(key=self.key, market=self.market, bucket=self.bucket, holding=self.holding)
        return d


def _iso(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "").replace("%", ""))
    except ValueError:
        return None


def _yes(v, default: bool) -> bool:
    if v is None or str(v).strip() == "":
        return default
    return str(v).strip().upper() in ("Y", "YES", "TRUE", "1")


def lock_file(path: Path) -> Path:
    return path.with_name("~$" + path.name)


class Book:
    """stocks.xlsx opened for reading (default) or writing (`write=True`)."""

    def __init__(self, path: Path = config.WORKBOOK, write: bool = False):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"{self.path} not found. Run ./setup.sh (or scripts/create_workbook.py).")
        if write and lock_file(self.path).exists():
            raise WorkbookOpenError(
                f"{self.path.name} is open in Excel (found {lock_file(self.path).name}). "
                "Save and close it, then re-run; nothing was written."
            )
        self.write = write
        self.wb = load_workbook(self.path)

    # --- helpers ---------------------------------------------------------

    @staticmethod
    def _cols(ws) -> dict[str, int]:
        return {str(c.value).strip(): c.column for c in ws[1] if c.value}

    def _sheet(self, name: str, columns: list[str]):
        ws = self.wb[name]
        cols = self._cols(ws)
        missing = [c for c in columns if c not in cols]
        for c in missing:  # older workbook: add new columns at the end
            idx = ws.max_column + 1
            ws.cell(row=1, column=idx, value=c).font = Font(bold=True)
            cols[c] = idx
        return ws, cols

    # --- Stocks ----------------------------------------------------------

    def stocks(self, settings: dict | None = None) -> list[StockRow]:
        settings = settings or self.settings()
        ws, c = self._sheet("Stocks", STOCK_COLUMNS)
        out = []
        for r in range(2, ws.max_row + 1):
            get = lambda name: ws.cell(row=r, column=c[name]).value  # noqa: E731
            ticker = get("Ticker")
            if not ticker or not str(ticker).strip():
                continue
            row = StockRow(
                row=r,
                ticker=str(ticker).strip().upper(),
                exchange=str(get("Exchange") or "").strip().upper(),
                company=get("Company"),
                status=str(get("Status") or "Watchlist").strip().title(),
                horizon=str(get("Investment Horizon") or "").strip(),
                qty=_num(get("Qty")),
                avg_price=_num(get("Avg Buy Price")),
                buy_date=_iso(get("Buy Date")),
                target_return=_num(get("Target Return %")),
                risk_tolerance=str(get("Risk Tolerance") or "Med").strip().title(),
                depth=(str(get("Depth")).strip().title() if get("Depth") else None) or settings["Default Depth"],
                second_opinion=_yes(get("2nd Opinion"), False),
                analyze=_yes(get("Analyze?"), True),
                notes=get("My Thesis / Notes"),
                last_analyzed=_iso(get("Last Analyzed")),
            )
            if row.exchange not in config.EXCHANGES:
                row.errors.append(f"Exchange {row.exchange!r} must be one of {', '.join(config.EXCHANGES)}")
            if row.horizon not in config.HORIZON_DAYS:
                row.errors.append(f"Investment Horizon {row.horizon!r} must be one of {', '.join(config.HORIZONS)}")
            if row.depth not in config.DEPTHS:
                row.errors.append(f"Depth {row.depth!r} must be one of {', '.join(config.DEPTHS)}")
            out.append(row)
        return out

    def update_stock(self, ticker: str, exchange: str, **fields) -> None:
        ws, c = self._sheet("Stocks", STOCK_COLUMNS)
        for r in range(2, ws.max_row + 1):
            t = str(ws.cell(row=r, column=c["Ticker"]).value or "").strip().upper()
            e = str(ws.cell(row=r, column=c["Exchange"]).value or "").strip().upper()
            if (t, e) != (ticker.upper(), exchange.upper()):
                continue
            for name, value in fields.items():
                cell = ws.cell(row=r, column=c[name])
                if name == "Company" and cell.value:
                    continue  # never overwrite what the user typed
                if name == "Last Report" and value:
                    cell.value, cell.hyperlink = "Open report", value
                    cell.font = Font(color="0563C1", underline="single")
                else:
                    cell.value = value
            return

    # --- Settings --------------------------------------------------------

    def settings(self) -> dict:
        s = dict(config.DEFAULT_SETTINGS)
        if "Settings" in self.wb.sheetnames:
            ws = self.wb["Settings"]
            for r in range(2, ws.max_row + 1):
                k, v = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
                if k and v not in (None, ""):
                    s[str(k).strip()] = v
        s["Default Depth"] = str(s["Default Depth"]).strip().title()
        s["Debate Rounds"] = int(_num(s["Debate Rounds"]) or 1)
        s["personas"] = {
            b: [p.strip().lower() for p in str(s[f"Personas ({b})"]).split(",") if p.strip()]
            for b in ("short", "medium", "long")
        }
        s["settle_days"] = {b: int(_num(s[f"Settle Days ({b})"]) or config.DEFAULT_SETTLE_DAYS[b])
                            for b in ("short", "medium", "long")}
        s["use_kite"] = _yes(s["Use Kite"], True)
        s["use_web_search"] = _yes(s["Use Web Search"], True)
        return s

    # --- Analysis Log ----------------------------------------------------

    def log_rows(self) -> list[dict]:
        ws, c = self._sheet("Analysis Log", LOG_COLUMNS)
        rows = []
        for r in range(2, ws.max_row + 1):
            d = {name: ws.cell(row=r, column=idx).value for name, idx in c.items()}
            if not d.get("Ticker"):
                continue
            d["_row"] = r
            for k in ("Date", "Data As Of", "Settled On"):
                d[k] = _iso(d.get(k))
            rows.append(d)
        return rows

    def history(self, ticker: str, exchange: str) -> list[dict]:
        """This stock's log rows, oldest first."""
        rows = [r for r in self.log_rows()
                if str(r["Ticker"]).upper() == ticker.upper() and str(r["Exchange"]).upper() == exchange.upper()]
        return sorted(rows, key=lambda r: r["Date"] or "")

    def upsert_log(self, entry: dict) -> int:
        """Write one run's row; a re-run on the same date replaces that day's row."""
        ws, c = self._sheet("Analysis Log", LOG_COLUMNS)
        target = None
        for r in range(2, ws.max_row + 1):
            key = tuple(_iso(ws.cell(row=r, column=c[k]).value) if k == "Date"
                        else str(ws.cell(row=r, column=c[k]).value or "").upper()
                        for k in ("Date", "Ticker", "Exchange"))
            if key == (entry["Date"], entry["Ticker"].upper(), entry["Exchange"].upper()):
                target = r
                break
        if target is None:
            target = ws.max_row + 1
            while target > 2 and not ws.cell(row=target - 1, column=c["Ticker"]).value:
                target -= 1
        for name, value in entry.items():
            if name not in c:
                continue
            cell = ws.cell(row=target, column=c[name])
            if name in ("Date", "Data As Of", "Settled On") and isinstance(value, str):
                value = date.fromisoformat(value)
                cell.number_format = "yyyy-mm-dd"
            if name == "Report" and value:
                cell.value, cell.hyperlink = "Open", value
                cell.font = Font(color="0563C1", underline="single")
            else:
                cell.value = value
            if name in ("Return", "Benchmark Return", "Alpha", "Confidence") and isinstance(value, float):
                cell.number_format = "0.0%"
        return target

    def settle(self, row: int, fields: dict) -> None:
        ws, c = self._sheet("Analysis Log", LOG_COLUMNS)
        for name, value in fields.items():
            cell = ws.cell(row=row, column=c[name])
            if name == "Settled On" and isinstance(value, str):
                value = date.fromisoformat(value)
                cell.number_format = "yyyy-mm-dd"
            cell.value = value
            if name in ("Return", "Benchmark Return", "Alpha"):
                cell.number_format = "0.0%"

    # --- Scorecard -------------------------------------------------------

    def rebuild_scorecard(self) -> None:
        settled = [r for r in self.log_rows() if r.get("Outcome") in ("Right", "Wrong")]
        if SCORECARD_SHEET in self.wb.sheetnames:
            del self.wb[SCORECARD_SHEET]
        ws = self.wb.create_sheet(SCORECARD_SHEET, index=2)
        ws.column_dimensions["A"].width = 26
        ws["A1"] = f"Scorecard: {len(settled)} settled calls (rebuilt {date.today().isoformat()})"
        ws["A1"].font = Font(bold=True, size=13)
        if not settled:
            ws["A3"] = "No calls settled yet. A call is scored once its settle window (Settings) has passed."
            return

        def bucket_of(r):
            try:
                return config.horizon_bucket(r.get("Horizon"))
            except ValueError:
                return "unknown"

        groupings = [
            ("By rating", lambda r: r.get("Rating") or "?"),
            ("By horizon", bucket_of),
            ("By stock", lambda r: f'{r["Ticker"]} ({r["Exchange"]})'),
        ]
        row = 3
        for title, keyfn in groupings:
            ws.cell(row=row, column=1, value=title).font = Font(bold=True)
            row += 1
            for i, h in enumerate(["Group", "Calls", "Right", "Hit rate", "Avg return", "Avg alpha"], 1):
                ws.cell(row=row, column=i, value=h).font = Font(bold=True)
            row += 1
            groups = defaultdict(list)
            for r in settled:
                groups[keyfn(r)].append(r)
            for name, rows in sorted(groups.items()):
                n = len(rows)
                right = sum(r["Outcome"] == "Right" for r in rows)
                avg = lambda k: sum(float(r[k] or 0) for r in rows) / n  # noqa: E731
                values = [name, n, right, right / n, avg("Return"), avg("Alpha")]
                for i, v in enumerate(values, 1):
                    cell = ws.cell(row=row, column=i, value=v)
                    if i >= 4:
                        cell.number_format = "0.0%"
                row += 1
            row += 1

    # --- Save --------------------------------------------------------------

    def save(self) -> None:
        if not self.write:
            raise RuntimeError("Book was opened read-only")
        if lock_file(self.path).exists():
            raise WorkbookOpenError(f"{self.path.name} was opened in Excel during the run; close it and re-run.")
        tmp = self.path.with_name(f".{self.path.stem}.tmp.xlsx")
        self.wb.save(tmp)
        os.replace(tmp, self.path)
